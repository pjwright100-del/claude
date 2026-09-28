"""Build the Chartis contractor invoice template (.xlsx, for import into Google Sheets).

Every date on the invoice is derived from the invoice number:
    period start = first period start + (invoice no. - 1) * 14 days
Only the cells on the Setup tab are edited each cycle.
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "1F4E79"
ACCENT = "5B9BD5"
TEXT = "404040"
FONT = "Arial"

wb = Workbook()

# ---------------------------------------------------------------- Setup tab
setup = wb.active
setup.title = "Setup"
setup.column_dimensions["A"].width = 34
setup.column_dimensions["B"].width = 30
setup.column_dimensions["C"].width = 60

input_fill = PatternFill("solid", fgColor="FFF2CC")
blue = Font(name=FONT, color="0000FF", size=11)
label = Font(name=FONT, size=11)
bold = Font(name=FONT, size=11, bold=True)

setup["A1"] = "Invoice Setup"
setup["A1"].font = Font(name=FONT, size=16, bold=True, color=NAVY)
setup["A2"] = "Edit only the yellow cells. The Invoice tab updates itself."
setup["A2"].font = Font(name=FONT, size=10, italic=True, color=TEXT)

rows = [
    # (row, label, value, number format, note, is_input)
    (4, "Invoice No.", 5, "0", "Each new invoice: add 1. All dates follow from this.", True),
    (5, "Week 1 hours", None, "General", "From Clockify (first week of the period)", True),
    (6, "Week 2 hours", None, "General", "From Clockify (second week of the period)", True),
    (7, "Other charge: description", "Anthropic Claude Pro Subscription",
     "@", "Optional extra line on the summary", True),
    (8, "Other charge: amount", None, '"$"#,##0.00',
     "Leave blank to hide the line", True),
    (10, "Hourly rate", 70, '"$"#,##0', "Rarely changes", True),
    (11, "First period start (invoice 0001)", date(2026, 7, 20), "m/d/yyyy",
     "Monday your first billing period began; rarely changes", True),
    (12, "Days per billing period", 14, "0", "Bi-weekly; rarely changes", True),
]
for r, lab, val, fmt, note, is_input in rows:
    setup.cell(r, 1, lab).font = label
    c = setup.cell(r, 2, val)
    c.number_format = fmt
    c.font = blue
    c.alignment = Alignment(horizontal="left")
    if is_input:
        c.fill = input_fill
    n = setup.cell(r, 3, note)
    n.font = Font(name=FONT, size=9, italic=True, color="808080")

setup["A14"] = "Calculated (do not edit)"
setup["A14"].font = bold
calc = [
    (15, "Period start", "=B11+(B4-1)*B12", "ddd m/d/yyyy"),
    (16, "Period end", "=B15+B12-1", "ddd m/d/yyyy"),
    (17, "Invoice date", "=B16+1", "ddd m/d/yyyy"),
]
for r, lab, f, fmt in calc:
    setup.cell(r, 1, lab).font = label
    c = setup.cell(r, 2, f)
    c.number_format = fmt
    c.font = Font(name=FONT, size=11)
    c.alignment = Alignment(horizontal="left")

dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1",
                    showErrorMessage=True, error="Invoice number must be 1 or more")
setup.add_data_validation(dv)
dv.add("B4")

# -------------------------------------------------------------- Invoice tab
inv = wb.create_sheet("Invoice")
inv.sheet_view.showGridLines = False
for col, w in zip("ABCD", (40, 14, 20, 20)):
    inv.column_dimensions[col].width = w

thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
navy_fill = PatternFill("solid", fgColor=NAVY)
f_text = Font(name=FONT, size=10, color=TEXT)
f_bold = Font(name=FONT, size=10, color=TEXT, bold=True)
f_head = Font(name=FONT, size=10, color="FFFFFF", bold=True)
f_section = Font(name=FONT, size=12, color=ACCENT, bold=True)


def put(ref, value, font=f_text, fmt=None, align="left", fill=None, border=None):
    c = inv[ref]
    c.value = value
    c.font = font
    c.alignment = Alignment(horizontal=align, vertical="center")
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = border
    return c


# Header: logo left, INVOICE right
inv.row_dimensions[1].height = 42
logo = Image("chartis_logo.png")
logo.width, logo.height = 170, 37
inv.add_image(logo, "A1")
put("D1", "INVOICE", Font(name=FONT, size=20, bold=True, color=NAVY), align="right")

put("A3", "Paul Wright", Font(name=FONT, size=14, bold=True, color=TEXT))
put("A4", "157 Hunter Ave, Fanwood, NJ 07023")
put("A5", '=TEXT(Setup!B10,"$#,##0")&"/hr"')

# Bill to / Invoice details
for ref, v in (("A7", "Bill to"), ("C7", "Invoice Details")):
    put(ref, v, f_head, fill=navy_fill)
for ref in ("B7", "D7"):
    inv[ref].fill = navy_fill
put("A8", "Cara Talmadge", f_bold)
put("A9", "Chartis Interactive LLC")
put("A10", "813-230-8376")
put("C8", "Invoice No:", f_bold)
put("D8", '=TEXT(Setup!B4,"0000")', align="right")
put("C9", "Invoice Date:", f_bold)
put("D9", "=Setup!B17", fmt="mmm d, yyyy", align="right")

# Billing period
put("A12", "Billing Period", f_head, fill=navy_fill)
inv["B12"].fill = navy_fill
put("C12", "Total Weeks", f_head, fill=navy_fill)
inv["D12"].fill = navy_fill
put("A13", '=TEXT(Setup!B15,"m/d/yyyy")&" – "&TEXT(Setup!B16,"m/d/yyyy")')
put("C13", '=ROUND(Setup!B12/7,0)&" weeks"')

# Clockify timesheet: date captions + empty areas for pasted screenshots
put("A15", "Clockify Timesheet", f_section)
put("A16", '=TEXT(Setup!B15,"m/d/yy")&" - "&TEXT(Setup!B15+6,"m/d/yy")', f_bold)
inv.row_dimensions[17].height = 70
put("A18", '=TEXT(Setup!B15+7,"m/d/yy")&" - "&TEXT(Setup!B16,"m/d/yy")', f_bold)
inv.row_dimensions[19].height = 70
for r in (17, 19):
    inv.merge_cells(f"A{r}:D{r}")
    put(f"A{r}", "Paste Clockify screenshot here (Insert > Image > Image over cells)",
        Font(name=FONT, size=9, italic=True, color="A6A6A6"), align="center")

# Invoice summary
put("A21", "Invoice Summary", f_section)
for col, h in zip("ABCD", ("By Week", "Hours", "Rate", "Line Total")):
    put(f"{col}22", h, f_head, fill=navy_fill, align="left" if col == "A" else "center",
        border=box)

money = '"$"#,##0.00'
put("A23", "=A16", border=box)
put("B23", "=N(Setup!B5)", fmt="General", align="center", border=box)
put("C23", "=Setup!B10", fmt='"$"#,##0', align="center", border=box)
put("D23", "=B23*C23", fmt=money, align="right", border=box)
put("A24", "=A18", border=box)
put("B24", "=N(Setup!B6)", fmt="General", align="center", border=box)
put("C24", "=Setup!B10", fmt='"$"#,##0', align="center", border=box)
put("D24", "=B24*C24", fmt=money, align="right", border=box)
# Optional other charge (blank when no amount)
put("A25", '=IF(N(Setup!B8)>0,Setup!B7,"")', border=box)
put("B25", None, border=box)
put("C25", None, border=box)
put("D25", '=IF(N(Setup!B8)>0,Setup!B8,"")', fmt=money, align="right", border=box)

put("A26", "Total Hours:", f_bold, border=box)
put("B26", "=B23+B24", f_bold, fmt="General", align="center", border=box)
put("C26", "Invoice Total:", f_bold, align="right", border=box)
put("D26", "=D23+D24+N(D25)", f_bold, fmt=money, align="right", border=box)

# Footer
inv.merge_cells("A28:D28")
put("A28", "PO Box 4509, Warren, NJ 07059   |   p. 813-230-8376   |   david@chartis.io",
    Font(name=FONT, size=9, color=TEXT), align="center")

# Print setup: one page wide, portrait, letter
inv.print_area = "A1:D28"
inv.page_setup.orientation = "portrait"
inv.page_setup.paperSize = inv.PAPERSIZE_LETTER
inv.page_setup.fitToWidth = 1
inv.page_setup.fitToHeight = 0
inv.sheet_properties.pageSetUpPr.fitToPage = True
inv.print_options.horizontalCentered = True

# Open on the Invoice tab
wb.move_sheet("Invoice", offset=-1)
wb.active = 0

wb.save("Chartis Invoice Template.xlsx")
print("saved")
